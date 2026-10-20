"""Extract plain text from uploaded proposal source documents (.docx / .pdf / .txt).

Used by the proposal module's document-ingestion flow (Phase 2c): the PM uploads
a finalized brief/notes document; we extract its text and hand it to the AI to
map into the nine-chapter proposal brief.
"""
from __future__ import annotations

import io


def extract_text_from_docx(content: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(content))
    parts = []
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            parts.append(p.text.strip())
    # Include table cell text too (briefs often use tables).
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text_from_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt.strip():
            parts.append(txt.strip())
    return "\n".join(parts)


def extract_text(filename: str, content: bytes) -> str:
    """Dispatch on file extension. Returns extracted plain text.

    Raises ValueError for unsupported types or empty extraction.
    """
    name = (filename or "").lower()
    if name.endswith(".docx"):
        text = extract_text_from_docx(content)
    elif name.endswith(".pdf"):
        text = extract_text_from_pdf(content)
    elif name.endswith(".txt") or name.endswith(".md"):
        text = content.decode("utf-8", errors="ignore")
    else:
        raise ValueError(
            "Unsupported file type. Please upload a .docx, .pdf, .txt or .md file."
        )

    text = (text or "").strip()
    if not text:
        raise ValueError(
            "Couldn't extract any text from that file. If it's a scanned/image PDF, "
            "paste the content into the builder chat instead."
        )
    # Cap very large docs to keep the AI prompt within limits.
    MAX_CHARS = 24000
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n[...truncated...]"
    return text
